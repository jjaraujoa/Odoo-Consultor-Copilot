#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
generar_acta_word.py
Genera la documentación formal de personalizaciones de Odoo en formato Word (.docx)
utilizando la plantilla estándar de entrega técnico-funcional y un archivo JSON estructurado.
"""

import sys
import os
import json
import argparse
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("ERROR: La librer\xc3\xada 'python-docx' no est\xc3\xa1 instalada en este entorno de Python.", file=sys.stderr)
    print("Ejecute: pip install python-docx", file=sys.stderr)
    sys.exit(1)

# Configurar stdout a UTF-8
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')


def build_word_document(json_data, template_path, output_path):
    template_path = Path(template_path).resolve()
    output_path = Path(output_path).resolve()

    if not template_path.exists():
        raise FileNotFoundError(f"No se encontr\xc3\xb3 la plantilla Word en: {template_path}")

    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document(str(template_path))

    # ==========================================
    # 1. TABLA 0: Ficha T\xc3\xa9cnica HU / Contexto
    # ==========================================
    t0_data = json_data.get("tabla_0_contexto", {})
    t0 = doc.tables[0]

    if "modulos" in t0_data:
        t0.rows[0].cells[1].text = str(t0_data["modulos"])
    if "version_odoo" in t0_data:
        t0.rows[1].cells[1].text = str(t0_data["version_odoo"])
    if "codigo_hu" in t0_data:
        t0.rows[2].cells[1].text = str(t0_data["codigo_hu"])
    if "codigo_gap_nombre" in t0_data:
        t0.rows[3].cells[1].text = str(t0_data["codigo_gap_nombre"])
    if "rol_solicitante" in t0_data:
        t0.rows[4].cells[1].text = str(t0_data["rol_solicitante"])
    if "objetivo_funcional" in t0_data:
        t0.rows[5].cells[1].text = str(t0_data["objetivo_funcional"])
    if "justificacion_criterio_exito" in t0_data:
        t0.rows[6].cells[1].text = str(t0_data["justificacion_criterio_exito"])

    for row in t0.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                for r in p.runs:
                    r.font.name = "Arial"
                    r.font.size = Pt(9.5)

    # ==========================================
    # 2. TABLA 1: Control de Entrega y Estimaci\xc3\xb3n
    # ==========================================
    t1_data = json_data.get("tabla_1_control", {})
    t1 = doc.tables[1]

    if "estimacion_horas" in t1_data:
        t1.rows[1].cells[0].text = str(t1_data["estimacion_horas"])
    if "equipo_especialidad" in t1_data:
        t1.rows[1].cells[1].text = str(t1_data["equipo_especialidad"])
    if "mes_ano" in t1_data:
        t1.rows[1].cells[2].text = str(t1_data["mes_ano"])

    if "etiqueta_entrega" in t1_data:
        t1.rows[2].cells[0].text = str(t1_data["etiqueta_entrega"])
    else:
        t1.rows[2].cells[0].text = "Fecha real de Entrega"

    if "consultor_responsable" in t1_data:
        t1.rows[2].cells[1].text = str(t1_data["consultor_responsable"])
    if "fecha_entrega" in t1_data:
        t1.rows[2].cells[2].text = str(t1_data["fecha_entrega"])

    for row in t1.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                for r in p.runs:
                    r.font.name = "Arial"
                    r.font.size = Pt(9.5)

    # ==========================================
    # 3. TABLA 2: Consideraciones y Exclusiones
    # ==========================================
    t2_data = json_data.get("tabla_2_consideraciones", [])
    t2 = doc.tables[2]
    t2_cell = t2.rows[0].cells[0]

    t2_p0 = t2_cell.paragraphs[0]
    t2_p0.paragraph_format.space_after = Pt(4)
    for r in t2_p0.runs:
        r.font.name = "Arial"

    for p in t2_cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)

    for item in t2_data:
        p = t2_cell.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Inches(0.2)
        run = p.add_run(item)
        run.font.name = "Arial"
        run.font.size = Pt(9)

    # ==========================================
    # 4. TABLA 3: Descripci\xc3\xb3n Funcional y T\xc3\xa9cnica
    # ==========================================
    t3_data = json_data.get("tabla_3_cuerpo", {})
    t3 = doc.tables[3]
    t3_cell = t3.rows[1].cells[0]

    p_first = t3_cell.paragraphs[0]
    p_first.text = ""
    for p in t3_cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)

    first_h1_used = False

    def add_h1(text):
        nonlocal first_h1_used
        if not first_h1_used:
            p = p_first
            first_h1_used = True
            for r in list(p.runs):
                p._p.remove(r._r)
        else:
            p = t3_cell.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(12)
        r.bold = True
        r.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        return p

    def add_h2(text):
        p = t3_cell.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10.5)
        r.bold = True
        r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        return p

    def add_p(text, bold=False, italic=False, space_after=3):
        p = t3_cell.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.bold = bold
        r.italic = italic
        return p

    def add_bullet(title, desc):
        p = t3_cell.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Inches(0.25)
        r_title = p.add_run("\xe2\u20ac\xa2 " + title + ": ")
        r_title.font.name = "Arial"
        r_title.font.size = Pt(9.5)
        r_title.bold = True
        r_desc = p.add_run(desc)
        r_desc.font.name = "Arial"
        r_desc.font.size = Pt(9.5)
        return p

    def add_code_block(text):
        p = t3_cell.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.3)
        r = p.add_run(text)
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(0x2A, 0x2A, 0x2A)
        return p

    def add_test_item(title, prep, res):
        p = t3_cell.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Inches(0.2)
        r1 = p.add_run(title + ": ")
        r1.bold = True
        r1.font.name = "Arial"
        r1.font.size = Pt(9.5)
        r2 = p.add_run("Condici\xc3\xb3n: " + prep + " -> Resultado esperado: " + res)
        r2.font.name = "Arial"
        r2.font.size = Pt(9)
        return p

    # --- 1. OBJETIVO GENERAL Y ALCANCE ---
    add_h1("1. OBJETIVO GENERAL Y ALCANCE DE LA SOLUCI\xc3\u201cN")
    sec1_text = t3_data.get("seccion_1_objetivo", "")
    add_p(sec1_text)

    # --- 2. MODELO DE DATOS Y CAMPOS CREADOS ---
    add_h1("2. MODELO DE DATOS Y CAMPOS CREADOS (ESTRUCTURA)")
    sec2_modelos = t3_data.get("seccion_2_modelo_datos", [])
    for idx, mod in enumerate(sec2_modelos, start=1):
        mod_nombre = mod.get("modelo", f"Modelo {idx}")
        add_h2(f"2.{idx} Modelo: {mod_nombre}")
        if "descripcion" in mod and mod["descripcion"]:
            add_p(mod["descripcion"])

        for campo in mod.get("campos", []):
            nombre_tec = campo.get("nombre_tecnico", "")
            etiqueta = campo.get("etiqueta", nombre_tec)
            tipo = campo.get("tipo", "Char")
            widget = campo.get("widget", "")
            notas = campo.get("notas", "")

            desc_parts = [f"Tipo {tipo}"]
            if widget:
                desc_parts.append(f"Widget '{widget}'")
            if notas:
                desc_parts.append(notas)

            desc = ". ".join(desc_parts) + "."
            bullet_title = f"{etiqueta} ({nombre_tec})" if nombre_tec != etiqueta else etiqueta
            add_bullet(bullet_title, desc)

        for nota in mod.get("notas_adicionales", []):
            if isinstance(nota, dict):
                add_bullet(nota.get("titulo", "Nota"), nota.get("descripcion", ""))
            else:
                add_p(str(nota))

    # --- 3. VISTAS Y ELEMENTOS DE INTERFAZ ---
    add_h1("3. VISTAS Y ELEMENTOS DE INTERFAZ DE USUARIO")
    sec3_vistas = t3_data.get("seccion_3_vistas_ui", [])
    for idx, v in enumerate(sec3_vistas, start=1):
        vista_nombre = v.get("vista", f"Vista {idx}")
        add_h2(f"3.{idx} {vista_nombre}")
        if "descripcion" in v and v["descripcion"]:
            add_p(v["descripcion"])
        for elem in v.get("elementos", []):
            add_bullet(elem.get("titulo", "Elemento"), elem.get("descripcion", ""))

    # --- 4. ACCIONES DEL SERVIDOR Y L\xc3\u201cGICA ---
    add_h1("4. ACCIONES DEL SERVIDOR Y LOGICA DE PROCESAMIENTO")
    sec4_acciones = t3_data.get("seccion_4_acciones_servidor", [])
    for idx, acc in enumerate(sec4_acciones, start=1):
        acc_nombre = acc.get("nombre", f"Acci\xc3\xb3n {idx}")
        add_h2(f"4.{idx} {acc_nombre}")
        if "modelo" in acc and acc["modelo"]:
            add_bullet("Modelo", acc["modelo"])
        if "tipo" in acc and acc["tipo"]:
            add_bullet("Tipo de acci\xc3\xb3n", acc["tipo"])

        logica = acc.get("logica", "")
        if isinstance(logica, list):
            add_bullet("L\xc3\xb3gica de negocio", "")
            for paso in logica:
                add_p(f"  \xe2\u20ac\xa2 {paso}", space_after=1)
        elif logica:
            add_bullet("L\xc3\xb3gica de negocio", logica)

        if "codigo" in acc and acc["codigo"]:
            add_bullet("Fragmento de c\xc3\xb3digo Python", "")
            add_code_block(acc["codigo"])

    # --- 5. REGLAS DE AUTOMATIZACI\xc3\u201cN ---
    add_h1("5. REGLAS DE AUTOMATIZACION (BASE.AUTOMATION)")
    sec5_reglas = t3_data.get("seccion_5_automatizaciones", [])
    for idx, reg in enumerate(sec5_reglas, start=1):
        reg_nombre = reg.get("nombre", f"Regla {idx}")
        add_h2(f"5.{idx} Regla: '{reg_nombre}'")
        if "modelo" in reg and reg["modelo"]:
            add_bullet("Modelo", reg["modelo"])
        if "disparador" in reg and reg["disparador"]:
            add_bullet("Disparador", reg["disparador"])
        if "campo_disparador" in reg and reg["campo_disparador"]:
            add_bullet("Campo disparador", reg["campo_disparador"])
        if "dominio_previo" in reg and reg["dominio_previo"]:
            add_bullet("Dominio previo", reg["dominio_previo"])
        if "dominio_posterior" in reg and reg["dominio_posterior"]:
            add_bullet("Dominio posterior", reg["dominio_posterior"])
        if "dominio_filtro" in reg and reg["dominio_filtro"]:
            add_bullet("Dominio de filtrado", reg["dominio_filtro"])
        if "accion" in reg and reg["accion"]:
            add_bullet("Acci\xc3\xb3n ejecutada", reg["accion"])
        if "estado" in reg and reg["estado"]:
            add_bullet("Estado configurado", reg["estado"])
        if "justificacion" in reg and reg["justificacion"]:
            add_bullet("Justificaci\xc3\xb3n funcional / Alcance", reg["justificacion"])
        if "funcionalidad" in reg and reg["funcionalidad"]:
            add_bullet("Funcionalidad", reg["funcionalidad"])

    # --- 6. PLANTILLA DE CORREO ELECTR\xc3\u201cNICO ---
    add_h1("6. PLANTILLA DE CORREO ELECTRONICO INSTITUCIONAL")
    sec6_plantillas = t3_data.get("seccion_6_plantillas_correo", [])
    if not sec6_plantillas:
        add_p("No aplica plantillas de correo para esta personalizaci\xc3\xb3n o se utilizan las nativas est\xc3\xa1ndar de Odoo.")
    else:
        for idx, plt in enumerate(sec6_plantillas, start=1):
            plt_nombre = plt.get("nombre", f"Plantilla {idx}")
            if len(sec6_plantillas) > 1:
                add_h2(f"6.{idx} Plantilla: '{plt_nombre}'")
            else:
                add_bullet("Nombre de la plantilla", plt_nombre)
            if "modelo" in plt and plt["modelo"]:
                add_bullet("Modelo aplicable", plt["modelo"])
            if "asunto" in plt and plt["asunto"]:
                add_bullet("Asunto", plt["asunto"])
            if "informes_adjuntos" in plt and plt["informes_adjuntos"]:
                add_bullet("Informe adjunto", plt["informes_adjuntos"])
            if "cuerpo_html" in plt and plt["cuerpo_html"]:
                add_bullet("Estructura del cuerpo HTML", plt["cuerpo_html"])

    # --- 7. CRITERIOS DE ACEPTACI\xc3\u201cN Y MATRIZ DE PRUEBAS ---
    add_h1("7. CRITERIOS DE ACEPTACI\xc3\u201cN Y MATRIZ DE PRUEBAS (DoD)")
    sec7_pruebas = t3_data.get("seccion_7_pruebas_dod", [])
    if not sec7_pruebas:
        add_p("Pruebas de aceptaci\xc3\xb3n verificadas conforme al est\xc3\xa1ndar de control de calidad.")
    else:
        for p_item in sec7_pruebas:
            p_titulo = p_item.get("titulo", "Prueba")
            p_cond = p_item.get("condicion", "N/A")
            p_res = p_item.get("resultado_esperado", "Conforme")
            add_test_item(p_titulo, p_cond, p_res)

    doc.save(str(output_path))
    print(f"\xc3\u2030XITO: Documento Word generado correctamente en: {output_path}")



def build_markdown_document(json_data, output_md_path):
    output_md_path = Path(output_md_path).resolve()
    output_md_path.parent.mkdir(parents=True, exist_ok=True)

    t0 = json_data.get("tabla_0_contexto", {})
    t1 = json_data.get("tabla_1_control", {})
    t2 = json_data.get("tabla_2_consideraciones", [])
    t3 = json_data.get("tabla_3_cuerpo", {})

    lines = []
    gap_title = t0.get("codigo_gap_nombre", "Personalizaci\u00f3n Odoo")
    lines.append(f"# Documentaci\u00f3n T\u00e9cnico-Funcional: {gap_title}\n")

    # Tabla 0
    lines.append("## Ficha T\u00e9cnica del Requerimiento\n")
    lines.append("| Campo | Detalle |")
    lines.append("| :--- | :--- |")
    lines.append(f"| **M\u00f3dulos por afectar** | {t0.get('modulos', 'N/D')} |")
    lines.append(f"| **Versi\u00f3n de Odoo** | {t0.get('version_odoo', 'N/D')} |")
    lines.append(f"| **Dependiente de HU** | {t0.get('codigo_hu', 'N/D')} |")
    lines.append(f"| **# GAP / Requerimiento** | {t0.get('codigo_gap_nombre', 'N/D')} |")
    lines.append(f"| **\u00bfQui\u00e9n? (Solicitante)** | {t0.get('rol_solicitante', 'N/D')} |")
    lines.append(f"| **\u00bfQuiero qu\u00e9? (Objetivo)** | {t0.get('objetivo_funcional', 'N/D')} |")
    lines.append(f"| **\u00bfPara qu\u00e9? (Criterio de \u00c9xito)** | {t0.get('justificacion_criterio_exito', 'N/D')} |\n")

    # Tabla 1
    lines.append("## Control de Entrega y Estimación\n")
    lines.append("| Métrica / Actividad | Especialidad / Responsable | Fecha / Periodo |")
    lines.append("| :--- | :--- | :--- |")
    val_est = t1.get("estimacion_horas", "Estimación: N/D")
    val_esp = t1.get("equipo_especialidad", "N/D")
    val_mes = t1.get("mes_ano", "N/D")
    lines.append(f"| {val_est} | {val_esp} | {val_mes} |")
    val_etiq = t1.get("etiqueta_entrega", "Fecha real de Entrega")
    val_resp = t1.get("consultor_responsable", "N/D")
    val_fent = t1.get("fecha_entrega", "N/D")
    lines.append(f"| {val_etiq} | {val_resp} | {val_fent} |\n")

    # Tabla 2
    lines.append("## Consideraciones, Reglas de Negocio y Exclusiones\n")
    for item in t2:
        lines.append(f"- {item}")
    lines.append("")

    # Tabla 3 - Secciones
    lines.append("## 1. Objetivo General y Alcance de la Soluci\u00f3n\n")
    lines.append(t3.get("seccion_1_objetivo", "") + "\n")

    lines.append("## 2. Modelo de Datos y Campos Creados (Estructura)\n")
    for idx, mod in enumerate(t3.get("seccion_2_modelo_datos", []), start=1):
        lines.append(f"### 2.{idx} Modelo: {mod.get('modelo', 'Modelo')}\n")
        if mod.get("descripcion"):
            lines.append(mod["descripcion"] + "\n")
        for campo in mod.get("campos", []):
            tec = campo.get("nombre_tecnico", "")
            lbl = campo.get("etiqueta", tec)
            tipo = campo.get("tipo", "Char")
            widget = campo.get("widget", "")
            notas = campo.get("notas", "")
            desc = f"Tipo {tipo}"
            if widget:
                desc += f", Widget '{widget}'"
            if notas:
                desc += f". {notas}"
            lines.append(f"- **{lbl}** (`{tec}`): {desc}.")
        for nota in mod.get("notas_adicionales", []):
            if isinstance(nota, dict):
                lines.append(f"- **{nota.get('titulo', 'Nota')}**: {nota.get('descripcion', '')}")
            else:
                lines.append(f"- {nota}")
        lines.append("")

    lines.append("## 3. Vistas y Elementos de Interfaz de Usuario\n")
    for idx, v in enumerate(t3.get("seccion_3_vistas_ui", []), start=1):
        lines.append(f"### 3.{idx} {v.get('vista', 'Vista')}\n")
        if v.get("descripcion"):
            lines.append(v["descripcion"] + "\n")
        for elem in v.get("elementos", []):
            lines.append(f"- **{elem.get('titulo', 'Elemento')}**: {elem.get('descripcion', '')}")
        lines.append("")

    lines.append("## 4. Acciones del Servidor y L\u00f3gica de Procesamiento\n")
    for idx, acc in enumerate(t3.get("seccion_4_acciones_servidor", []), start=1):
        acc_nom = acc.get("nombre", "Acción")
        lines.append(f"### 4.{idx} {acc_nom}\n")
        if acc.get("modelo"):
            lines.append(f"- **Modelo**: `{acc['modelo']}`")
        if acc.get("tipo"):
            lines.append(f"- **Tipo de acci\u00f3n**: {acc['tipo']}")
        logica = acc.get("logica", "")
        if isinstance(logica, list):
            lines.append("- **L\u00f3gica de negocio**:")
            for p in logica:
                lines.append(f"  - {p}")
        elif logica:
            lines.append(f"- **L\u00f3gica de negocio**: {logica}")
        if acc.get("codigo"):
            lines.append("\n```python")
            lines.append(acc["codigo"])
            lines.append("```\n")
        else:
            lines.append("")

    lines.append("## 5. Reglas de Automatizaci\u00f3n (base.automation)\n")
    for idx, reg in enumerate(t3.get("seccion_5_automatizaciones", []), start=1):
        lines.append(f"### 5.{idx} Regla: '{reg.get('nombre', 'Regla')}'\n")
        if reg.get("modelo"):
            lines.append(f"- **Modelo**: `{reg['modelo']}`")
        if reg.get("disparador"):
            lines.append(f"- **Disparador**: {reg['disparador']}")
        if reg.get("campo_disparador"):
            lines.append(f"- **Campo disparador**: `{reg['campo_disparador']}`")
        if reg.get("dominio_posterior"):
            lines.append(f"- **Dominio posterior**: `{reg['dominio_posterior']}`")
        if reg.get("accion"):
            lines.append(f"- **Acci\u00f3n ejecutada**: {reg['accion']}")
        if reg.get("estado"):
            lines.append(f"- **Estado configurado**: `{reg['estado']}`")
        if reg.get("justificacion"):
            lines.append(f"- **Justificaci\u00f3n**: {reg['justificacion']}")
        lines.append("")

    lines.append("## 6. Plantilla de Correo Electr\u00f3nico Institucional\n")
    plantillas = t3.get("seccion_6_plantillas_correo", [])
    if not plantillas:
        lines.append("No aplica plantillas de correo para esta personalizaci\u00f3n.\n")
    else:
        for idx, plt in enumerate(plantillas, start=1):
            lines.append(f"### 6.{idx} Plantilla: '{plt.get('nombre', 'Plantilla')}'\n")
            if plt.get("modelo"):
                lines.append(f"- **Modelo aplicable**: `{plt['modelo']}`")
            if plt.get("asunto"):
                lines.append(f"- **Asunto**: `{plt['asunto']}`")
            if plt.get("informes_adjuntos"):
                lines.append(f"- **Informe adjunto**: `{plt['informes_adjuntos']}`")
            if plt.get("cuerpo_html"):
                lines.append(f"- **Estructura HTML**: {plt['cuerpo_html']}")
            lines.append("")

    lines.append("## 7. Criterios de Aceptaci\u00f3n y Matriz de Pruebas (DoD)\n")
    for p_item in t3.get("seccion_7_pruebas_dod", []):
        lines.append(f"- **{p_item.get('titulo', 'Prueba')}**:")
        lines.append(f"  - **Condici\u00f3n**: {p_item.get('condicion', 'N/D')}")
        lines.append(f"  - **Resultado esperado**: {p_item.get('resultado_esperado', 'Conforme')}")
    lines.append("")

    with open(output_md_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print(f"\u00c9XITO: Documento Markdown generado correctamente en: {output_md_path}")


def main():
    parser = argparse.ArgumentParser(description="Generador de Actas de Entrega Word (.docx) para Odoo")
    parser.add_argument("--json", required=True, help="Ruta al archivo JSON con los datos estructurados")
    parser.add_argument("--template", default=None, help="Ruta a la plantilla Word (.docx) base (opcional)")
    parser.add_argument("--output", required=True, help="Ruta de destino para el archivo .docx generado")
    parser.add_argument("--markdown", default=None, help="Ruta de destino para el archivo .md generado (opcional)")

    args = parser.parse_args()

    json_path = Path(args.json).resolve()
    if not json_path.exists():
        print(f"ERROR: No se encontr\xc3\xb3 el archivo JSON en: {json_path}", file=sys.stderr)
        sys.exit(1)

    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    if args.template:
        template_path = Path(args.template).resolve()
    else:
        script_dir = Path(__file__).parent.resolve()
        std_tpl = script_dir.parent / "resources" / "plantilla_acta_entrega_estandar.docx"
        legacy_tpl = script_dir.parent / "resources" / "plantilla_acta_entrega.docx"
        template_path = std_tpl if std_tpl.exists() else legacy_tpl

    build_word_document(data, template_path, args.output)
    if args.markdown:
        build_markdown_document(data, args.markdown)


if __name__ == "__main__":
    main()
