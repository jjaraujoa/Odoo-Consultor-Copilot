#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
scripts/construir_matriz.py - Compilador de Matriz de GAPs de Odoo (Estándar de Consultoría)
Toma un archivo JSON estructurado con requerimientos evaluados y genera el libro Excel (.xlsx).

En sistemas Windows con Microsoft Excel, delega automáticamente en construir_matriz.ps1
para utilizar el motor COM nativo de Microsoft Excel, garantizando:
- Casillas de verificación interactivas nativas de Excel 365 (sin degradar a VERDADERO/FALSO).
- Fórmulas dinámicas vinculadas y casillas en el Consolidado.
- Preservación íntegra de estilos corporativos, logos y alturas de fila.
"""

import os
import sys
import json
import subprocess
from datetime import datetime

# Salida UTF-8
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

def compilar_con_powershell_com(json_path, template_path=None, output_path=None):
    """Ejecuta el compilador nativo PowerShell con Excel COM."""
    script_dir = os.path.dirname(os.path.abspath(__file__))
    ps1_path = os.path.join(script_dir, "construir_matriz.ps1")

    if not os.path.exists(ps1_path):
        return False, f"No se encontró el script {ps1_path}"

    cmd = [
        "powershell",
        "-NoProfile",
        "-ExecutionPolicy", "Bypass",
        "-File", ps1_path,
        "-InputJson", os.path.abspath(json_path)
    ]
    if template_path:
        cmd.extend(["-TemplatePath", os.path.abspath(template_path)])
    if output_path:
        cmd.extend(["-OutputPath", os.path.abspath(output_path)])

    res = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace")
    if res.returncode == 0:
        return True, res.stdout
    else:
        return False, f"PowerShell COM Error:\nSTDOUT:\n{res.stdout}\nSTDERR:\n{res.stderr}"

def compilar_fallback_openpyxl(json_path, template_path=None, output_path=None):
    """Compilador de reserva usando openpyxl si COM no está disponible."""
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    cliente = data.get("cliente", "CLIENTE")
    fecha_str = data.get("fecha", datetime.now().strftime("%Y-%m-%d"))
    version = data.get("version", "1")
    hojas_data = data.get("hojas", {})

    if not template_path:
        base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        template_path = os.path.join(base_dir, "assets", "FormatoMatrizDeGaps-Estandar-V1.xlsx")

    if not os.path.exists(template_path):
        raise FileNotFoundError(f"No se encontró la plantilla Excel: {template_path}")

    if not output_path:
        cli_clean = "".join(c for c in cliente if c.isalnum() or c in ('_', '-'))
        f_stamp = datetime.now().strftime("%Y%m%d")
        out_dir = os.path.join(os.getcwd(), "outputs")
        os.makedirs(out_dir, exist_ok=True)
        output_path = os.path.join(out_dir, f"FormatoMatrizDeGaps-{cli_clean}-{f_stamp}-V{version}.xlsx")

    wb = openpyxl.load_workbook(template_path)

    consecutivo_global = 1
    consolidado_row = 5
    ws_cons = wb["Consolidado"] if "Consolidado" in wb.sheetnames else None

    # Limpiar consolidado
    if ws_cons:
        for r in range(5, max(ws_cons.max_row + 1, 50)):
            for c in range(1, 18):
                ws_cons.cell(r, c).value = None
        ws_cons["P2"].value = f"V{version}"
        ws_cons["P3"].value = fecha_str

    # Limpiar hojas preexistentes no utilizadas en este conjunto de requerimientos
    for sname in list(wb.sheetnames):
        if sname != "Consolidado" and sname not in hojas_data:
            ws_unused = wb[sname]
            for r in range(5, max(ws_unused.max_row + 1, 50)):
                for c in range(1, 18):
                    ws_unused.cell(r, c).value = None

    for sname, reqs in hojas_data.items():
        if sname == "Consolidado":
            continue

        if sname in wb.sheetnames:
            ws = wb[sname]
        else:
            ref = wb.worksheets[0]
            ws = wb.copy_worksheet(ref)
            ws.title = sname

        ws["P2"].value = f"V{version}"
        ws["P3"].value = fecha_str

        # Limpiar filas previas
        for r in range(5, max(ws.max_row + 1, 25)):
            for c in range(1, 18):
                ws.cell(r, c).value = None

        for i, req in enumerate(reqs):
            r = 5 + i
            req_nombre = req.get("requerimiento", "")
            req_cia = req.get("compania", cliente)
            req_fecha = req.get("fecha", fecha_str)
            req_mod = req.get("odoo_modulo", sname)
            req_f = float(req.get("porcentaje_ajuste_nativo", 1.0))

            req_g = bool(req.get("localizacion_regional", False))
            req_h = bool(req.get("desarrollo", False))
            req_i = bool(req.get("studio", False))
            req_j = bool(req.get("permisos", False))
            req_k = bool(req.get("accion_automatica", False))

            req_l = req.get("como_cumple_odoo", "")
            req_m = req.get("observaciones", "")
            req_o = req.get("prioridad", "Media")
            req_p = req.get("observaciones_desarrollo", "") if req_h else ""
            req_n = "No" if req_f >= 1.0 else "Si"

            if req_f >= 1.0:
                req_g = req_h = req_i = req_j = req_k = False
                req_n = "No"
                req_p = ""

            ws.cell(r, 1, i + 1)
            ws.cell(r, 2, req_nombre)
            ws.cell(r, 3, req_cia)
            ws.cell(r, 4, req_fecha)
            ws.cell(r, 5, req_mod)
            c_f = ws.cell(r, 6, req_f)
            c_f.number_format = "0%"

            ws.cell(r, 7, req_g)
            ws.cell(r, 8, req_h)
            ws.cell(r, 9, req_i)
            ws.cell(r, 10, req_j)
            ws.cell(r, 11, req_k)

            ws.cell(r, 12, req_l)
            ws.cell(r, 13, req_m)
            ws.cell(r, 14, req_n)
            ws.cell(r, 15, req_o)
            ws.cell(r, 16, req_p)

            if ws_cons:
                cr = consolidado_row
                ws_cons.cell(cr, 1, consecutivo_global)
                ws_cons.cell(cr, 2, f"='{sname}'!B{r}")
                ws_cons.cell(cr, 3, f"='{sname}'!C{r}")
                ws_cons.cell(cr, 4, f"='{sname}'!D{r}")
                ws_cons.cell(cr, 5, f"='{sname}'!E{r}")
                c_fc = ws_cons.cell(cr, 6, f"='{sname}'!F{r}")
                c_fc.number_format = "0%"

                ws_cons.cell(cr, 7, f"='{sname}'!G{r}")
                ws_cons.cell(cr, 8, f"='{sname}'!H{r}")
                ws_cons.cell(cr, 9, f"='{sname}'!I{r}")
                ws_cons.cell(cr, 10, f"='{sname}'!J{r}")
                ws_cons.cell(cr, 11, f"='{sname}'!K{r}")

                ws_cons.cell(cr, 12, f"='{sname}'!L{r}")
                ws_cons.cell(cr, 13, f"='{sname}'!M{r}")
                ws_cons.cell(cr, 14, f"='{sname}'!N{r}")
                ws_cons.cell(cr, 15, f"='{sname}'!O{r}")
                ws_cons.cell(cr, 16, f"='{sname}'!P{r}")

                consecutivo_global += 1
                consolidado_row += 1

    wb.save(output_path)
    return True, f"Compilado con openpyxl en: {output_path}"

def main():
    import argparse
    parser = argparse.ArgumentParser(description="Compilador de Matriz de GAPs de Odoo")
    parser.add_argument("json_path", help="Ruta al archivo JSON con los requerimientos")
    parser.add_argument("pos_template", nargs="?", default=None, help="Ruta a plantilla (posicional)")
    parser.add_argument("pos_output", nargs="?", default=None, help="Ruta de salida (posicional)")
    parser.add_argument("--template", dest="flag_template", default=None, help="Ruta a la plantilla Excel base")
    parser.add_argument("--output", dest="flag_output", default=None, help="Ruta de destino para el archivo .xlsx generado")

    args = parser.parse_args()
    json_path = args.json_path
    template_path = args.flag_template or args.pos_template
    output_path = args.flag_output or args.pos_output

    # Intentar compilación COM en Windows para casillas interactivas
    if os.name == 'nt':
        success, msg = compilar_con_powershell_com(json_path, template_path, output_path)
        if success:
            print(msg)
            sys.exit(0)
        else:
            print(f"Aviso COM no disponible, usando fallback openpyxl: {msg}")

    success, msg = compilar_fallback_openpyxl(json_path, template_path, output_path)
    print(msg)
    sys.exit(0 if success else 1)

if __name__ == "__main__":
    main()
